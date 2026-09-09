"""Tests for the Spec Analyzer's Word report writer."""

import docx

from qa_agents.spec_analyzer.models import AnalysisResult, Entity
from qa_agents.spec_analyzer.word_writer import write_report


def _sample_result() -> AnalysisResult:
    return AnalysisResult(
        business_rules=["חוק עסקי ראשון.", "חוק עסקי שני."],
        entities=[Entity(name="הזמנה (demo_order)", fields=["demo_orderstatus", "demo_supplierid"])],
        gaps=["לא מוגדרת התנהגות לקובץ לא תקין."],
        recommendations=["להבהיר את הטיפול בקובץ לא תקין."],
    )


def _paragraph_texts(path):
    document = docx.Document(path)
    return [(p.style.name, p.text) for p in document.paragraphs]


def test_writes_title_and_scope_for_a_tag(tmp_path):
    out = write_report(_sample_result(), tmp_path / "analysis.docx", tag="40100")
    texts = _paragraph_texts(out)
    assert ("Heading 1", "דוח ניתוח אפיון") in texts
    assert any("40100" in text for _, text in texts)


def test_writes_scope_for_whole_spec(tmp_path):
    out = write_report(_sample_result(), tmp_path / "analysis.docx", tag=None)
    texts = _paragraph_texts(out)
    assert any("כלל האפיון" in text for _, text in texts)


def test_writes_scope_for_image_without_tag(tmp_path):
    out = write_report(
        _sample_result(), tmp_path / "analysis.docx", tag=None, images_attached=True
    )
    texts = [t for _, t in _paragraph_texts(out)]
    assert any("תמונה מצורפת" in t for t in texts)
    assert not any(t == "מיקוד: כלל האפיון" for t in texts)


def test_tag_takes_priority_over_image_scope_in_header(tmp_path):
    out = write_report(
        _sample_result(), tmp_path / "analysis.docx", tag="40100", images_attached=True
    )
    texts = [t for _, t in _paragraph_texts(out)]
    assert any("40100" in t for t in texts)
    assert not any("תמונה מצורפת" in t for t in texts)


def test_writes_all_sections_and_bullets(tmp_path):
    out = write_report(_sample_result(), tmp_path / "analysis.docx")
    texts = _paragraph_texts(out)

    section_titles = [t for style, t in texts if style == "Heading 2"]
    assert section_titles == [
        "חוקים עסקיים שזוהו",
        "ישויות ושדות שאותרו",
        "פערים ואי-בהירויות באפיון",
        "המלצות להשלמה לפני תחילת בדיקות",
    ]

    bullets = [t for style, t in texts if style == "List Bullet"]
    assert "חוק עסקי ראשון." in bullets
    assert "demo_orderstatus" in bullets
    assert "לא מוגדרת התנהגות לקובץ לא תקין." in bullets
    assert "להבהיר את הטיפול בקובץ לא תקין." in bullets

    entity_headings = [t for style, t in texts if style == "Heading 3"]
    assert entity_headings == ["הזמנה (demo_order)"]


def test_empty_sections_show_placeholder(tmp_path):
    out = write_report(AnalysisResult(), tmp_path / "empty.docx")
    texts = [t for _, t in _paragraph_texts(out)]
    assert texts.count("לא זוהה תוכן רלוונטי.") == 4  # rules, entities, gaps, recommendations


def test_paragraphs_are_rtl(tmp_path):
    out = write_report(_sample_result(), tmp_path / "analysis.docx")
    document = docx.Document(out)
    body_paragraphs = [p for p in document.paragraphs if p.text]
    assert body_paragraphs  # sanity: there is content
    for p in body_paragraphs:
        p_pr = p._p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
        assert p_pr is not None
        bidi = p_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi")
        assert bidi is not None, f"paragraph {p.text!r} is missing the RTL (w:bidi) flag"
