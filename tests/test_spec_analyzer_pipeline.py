"""End-to-end test: spec document → analysis Word report, with a mocked LLM."""

import json

import docx as docx_lib

from qa_agents.spec_analyzer.agent import SpecAnalyzerAgent
from qa_agents.spec_analyzer.pipeline import default_output_name, generate_analysis

_CANNED = json.dumps(
    {
        "business_rules": ["בקשה מאושרת רק אם הלקוח פעיל וקיים סכום תקין."],
        "entities": [{"name": "בקשת החזר (req_refund)", "fields": ["req_status", "req_amount"]}],
        "gaps": ["לא ברור מה קורה כשהלקוח חסום באמצע טיפול בבקשה."],
        "recommendations": ["להבהיר את הטיפול בלקוח שנחסם באמצע תהליך."],
    },
    ensure_ascii=False,
)


def test_default_output_name():
    assert default_output_name("40100") == "output/ANALYSIS_40100.docx"
    assert default_output_name("full_spec") == "output/ANALYSIS_full_spec.docx"


def test_generate_analysis_from_docx_produces_report(tmp_path):
    document = docx_lib.Document()
    document.add_paragraph("תיוג 40100: אישור בקשת החזר כספי ללקוח.")
    spec = tmp_path / "spec.docx"
    document.save(spec)

    agent = SpecAnalyzerAgent(completer=lambda system, user: _CANNED)
    out = generate_analysis(spec, tag="40100", output_path=tmp_path / "analysis.docx", agent=agent)

    report = docx_lib.Document(out)
    texts = [p.text for p in report.paragraphs]
    assert any("40100" in t for t in texts)
    assert "בקשה מאושרת רק אם הלקוח פעיל וקיים סכום תקין." in texts
    assert "req_status" in texts
