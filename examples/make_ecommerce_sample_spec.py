"""Generate a richer, fully fictional English e-commerce demo specification.

Run:  python examples/make_ecommerce_sample_spec.py
A realistic-looking process spec for the ecommerce-english profile — invented
data, no real organization. Mirrors make_sample_spec.py's structure/complexity
so the two profiles can be exercised side by side.
"""

from pathlib import Path

import docx


def build() -> Path:
    d = docx.Document()
    d.add_heading("Process Specification — Coupon Code Redemption", level=1)
    d.add_paragraph("Task tag: 50200")
    d.add_paragraph(
        "This document defines the process of redeeming a discount coupon code "
        "at checkout, including eligibility conditions, field updates, and exceptions."
    )

    d.add_heading("Entities and fields", level=2)

    d.add_paragraph('Entity "Coupon" (promo_coupon):')
    t1 = d.add_table(rows=1, cols=3)
    t1.rows[0].cells[0].text = "Field name"
    t1.rows[0].cells[1].text = "Technical name"
    t1.rows[0].cells[2].text = "Notes"
    for name, tech, note in [
        ("Coupon status", "promo_status", "Code: 1=active, 2=expired, 3=disabled"),
        ("Discount type", "promo_discounttype", "Code: 1=percentage, 2=fixed amount"),
        ("Discount value", "promo_discountvalue", "Numeric — percentage or currency amount"),
        ("Minimum order amount", "promo_minorder", "Numeric — order must reach this to qualify"),
        ("Usage limit", "promo_usagelimit", "Numeric — max number of redemptions allowed"),
        ("Times used", "promo_timesused", "Numeric — incremented on each successful redemption"),
        ("Expiry date", "promo_expirydate", "Coupon is invalid strictly after this date"),
    ]:
        c = t1.add_row().cells
        c[0].text, c[1].text, c[2].text = name, tech, note

    d.add_paragraph('Entity "Order" (shop_order):')
    t2 = d.add_table(rows=1, cols=3)
    t2.rows[0].cells[0].text = "Field name"
    t2.rows[0].cells[1].text = "Technical name"
    t2.rows[0].cells[2].text = "Notes"
    for name, tech, note in [
        ("Order status", "order_status", "Code: 1=pending, 2=paid, 3=cancelled"),
        ("Order total", "order_total", "Numeric — pre-discount order amount"),
        ("Customer ID", "order_customerid", "Foreign key to the Customer entity (shop_customer)"),
        ("Applied coupon code", "order_couponcode", "Set only on successful redemption"),
        ("Discount applied", "order_discountapplied", "Numeric — resulting discount amount; 0 if none"),
    ]:
        c = t2.add_row().cells
        c[0].text, c[1].text, c[2].text = name, tech, note

    d.add_paragraph('Entity "Customer" (shop_customer):')
    t3 = d.add_table(rows=1, cols=3)
    t3.rows[0].cells[0].text = "Field name"
    t3.rows[0].cells[1].text = "Technical name"
    t3.rows[0].cells[2].text = "Notes"
    for name, tech, note in [
        ("Customer status", "cust_status", "Code: 1=active, 2=blocked"),
    ]:
        c = t3.add_row().cells
        c[0].text, c[1].text, c[2].text = name, tech, note

    d.add_heading("Business rules", level=2)
    for rule in [
        "A coupon may be redeemed only if: promo_status = 1 (active), today's date is on or "
        "before promo_expirydate, promo_timesused is strictly less than promo_usagelimit, and "
        "order_total is greater than or equal to promo_minorder.",
        "On successful redemption: order_couponcode is set to the coupon's code, "
        "order_discountapplied is calculated from promo_discounttype and promo_discountvalue "
        "(percentage of order_total, or a fixed amount), and promo_timesused is incremented by 1.",
        "If the customer is blocked (cust_status = 2), redemption is rejected regardless of "
        "coupon validity — order_couponcode and order_discountapplied remain unchanged.",
        "If promo_timesused has already reached promo_usagelimit, redemption is rejected even "
        "if the coupon is otherwise active and not expired.",
        "An order that already has a coupon applied (order_couponcode is not empty) cannot have "
        "a second coupon applied — the redemption request is rejected as a no-op.",
        "A cancelled order (order_status = 3) cannot have a coupon applied or removed.",
    ]:
        d.add_paragraph(rule, style="List Bullet")

    d.add_heading("Interfaces", level=2)
    d.add_paragraph(
        "On successful coupon redemption, a notification is sent to the loyalty/marketing "
        "system. Sending the notification itself is out of scope for this test, but a log line "
        "confirming the notification was sent must be verified."
    )

    out = Path(__file__).with_name("ecommerce_sample_spec.docx")
    d.save(out)
    return out


if __name__ == "__main__":
    build()
    print("Wrote ecommerce_sample_spec.docx")
