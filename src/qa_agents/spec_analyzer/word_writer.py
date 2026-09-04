"""Build the analysis Word report from an ``AnalysisResult``.

The engine — not the LLM — writes the file, matching the STD Generator's
principle: the model produces structured data, plain code turns it into the
deliverable. python-docx has no high-level "RTL paragraph" property, so it is
set via the underlying OOXML ``w:bidi`` element.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .models import AnalysisResult

_NO_CONTENT = "לא זוהה תוכן רלוונטי."


def _rtl(paragraph: Paragraph) -> Paragraph:
    """Mark a paragraph right-to-left (alignment + the OOXML bidi flag)."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(p_pr.makeelement(qn("w:bidi"), {}))
    return paragraph


def _heading(document: docx.document.Document, text: str, level: int) -> None:
    _rtl(document.add_heading(text, level=level))


def _bullets(document: docx.document.Document, items: list[str]) -> None:
    if not items:
        _rtl(document.add_paragraph(_NO_CONTENT))
        return
    for item in items:
        _rtl(document.add_paragraph(item, style="List Bullet"))


def write_report(result: AnalysisResult, path: str | Path, tag: str | None = None) -> Path:
    """Write ``result`` to a ``.docx`` file and return its path."""
    document = docx.Document()

    _heading(document, "דוח ניתוח אפיון", level=1)
    _rtl(document.add_paragraph(f'מיקוד: תיוג {tag}' if tag else "מיקוד: כלל האפיון"))

    _heading(document, "חוקים עסקיים שזוהו", level=2)
    _bullets(document, result.business_rules)

    _heading(document, "ישויות ושדות שאותרו", level=2)
    if not result.entities:
        _rtl(document.add_paragraph(_NO_CONTENT))
    for entity in result.entities:
        _heading(document, entity.name, level=3)
        _bullets(document, entity.fields)

    _heading(document, "פערים ואי-בהירויות באפיון", level=2)
    _bullets(document, result.gaps)

    _heading(document, "המלצות להשלמה לפני תחילת בדיקות", level=2)
    _bullets(document, result.recommendations)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    return out
